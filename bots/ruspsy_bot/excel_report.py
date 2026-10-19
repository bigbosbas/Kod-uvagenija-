import io
import logging
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
import database as db

log = logging.getLogger(__name__)

FONT = "Arial"
CLR_DARK = "1F3864"
CLR_ACCENT = "2E75B6"
CLR_LIGHT = "DEEBF7"

INT = '#,##0;[Red](#,##0);"—"'
MONEY = '#,##0 ₽;[Red](#,##0 ₽);"—"'
PERCENT = '0.0%;[Red]-0.0%;"—"'
DATE = "mm-dd-yy"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SEG_TITLES = dict(config.PRODUCT_TITLES)
SEG_TITLES.update({None: "не выбрал", "": "не выбрал"})
SRC_TITLES = {"ig": "Instagram", "yt": "YouTube", "direct": "Прямой заход"}


def _title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=CLR_DARK)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")


def _header(ws, row, columns, widths):
    for idx, (title, width) in enumerate(zip(columns, widths), start=1):
        cell = ws.cell(row=row, column=idx, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", start_color=CLR_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[row].height = 30


def _rows(ws, start_row, data, formats=None):
    for r, row in enumerate(data, start=start_row):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if formats and formats.get(c):
                cell.number_format = formats[c]
    return start_row + len(data)


def _dt(value, shift=True):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value).replace(microsecond=0)
    except ValueError:
        return value
    return dt + timedelta(hours=config.TZ_HOURS) if shift else dt


def _sheet_funnel(wb):
    ws = wb.create_sheet("Воронка")
    _title(ws, "Воронка", "Сколько человек дошло до каждого шага")
    _header(ws, 4, ["Шаг", "Человек", "От запустивших", "От предыдущего шага"], [34, 14, 18, 22])

    steps = db.funnel()
    first = steps[0][1] or 0
    data = []
    for i, (name, count) in enumerate(steps):
        prev = steps[i - 1][1] if i else count
        data.append([name, count,
                     (count / first) if first else 0,
                     (count / prev) if prev else 0])
    end = _rows(ws, 5, data, {2: INT, 3: PERCENT, 4: PERCENT})
    ws.freeze_panes = "A5"

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Воронка"
    chart.height, chart.width = 8, 18
    chart.add_data(Reference(ws, min_col=2, min_row=4, max_row=end - 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=end - 1))
    ws.add_chart(chart, "F4")
    return end


def _sheet_users(wb):
    ws = wb.create_sheet("Пользователи")
    _title(ws, "Пользователи", "Все, кто запускал бота")
    _header(ws, 4, ["ID", "Username", "Имя", "Источник", "Согласие", "Направление",
                    "Купил", "Шаг дожима", "Заблокировал", "Пришёл"],
            [14, 18, 18, 14, 11, 16, 18, 12, 13, 18])

    data = []
    for u in db.all_users():
        data.append([
            u["user_id"],
            f"@{u['username']}" if u["username"] else "",
            u["first_name"] or "",
            SRC_TITLES.get(u["source"], u["source"] or ""),
            "да" if u["consent"] else "нет",
            SEG_TITLES.get(u["segment"], u["segment"] or "не выбрал"),
            ", ".join(SEG_TITLES.get(p, p) for p in (u["bought"] or "").split(",") if p),
            u["drip_step"] if u["drip_step"] is not None else "",
            "да" if u["blocked"] else "",
            _dt(u["created_at"]),
        ])
    _rows(ws, 5, data, {1: INT, 10: DATE})
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{4 + len(data)}"
    return len(data)


def _sheet_payments(wb):
    ws = wb.create_sheet("Платежи")
    _title(ws, "Платежи", "Все оплаты: сумма, покупатель, товар")
    _header(ws, 4, ["Когда", "Заказ", "Товар", "Сумма", "ID", "Username", "Имя",
                    "E-mail", "Источник"],
            [18, 9, 24, 12, 14, 18, 16, 26, 16])
    rows = db.paid_orders()
    data = []
    for o in rows:
        from_site = o["user_id"] is None
        data.append([
            _dt(o["paid_at"]), o["order_id"],
            SEG_TITLES.get(o["product"], o["product"]), o["amount_rub"] or 0,
            o["user_id"] or "", f"@{o['username']}" if o["username"] else "",
            o["first_name"] or "", o["email"] or "",
            "Сайт (витрина)" if from_site
            else SRC_TITLES.get(o["source"], o["source"] or "")])
    end = _rows(ws, 5, data, {1: DATE, 2: INT, 4: MONEY, 5: INT})

    total = sum(o["amount_rub"] or 0 for o in rows)
    ws.cell(row=end, column=3, value="ИТОГО").font = Font(name=FONT, bold=True, size=10)
    tc = ws.cell(row=end, column=4, value=total)
    tc.font = Font(name=FONT, bold=True, size=11, color=CLR_DARK)
    tc.number_format = MONEY
    for col in range(1, 10):
        ws.cell(row=end, column=col).fill = PatternFill("solid", start_color=CLR_LIGHT)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{4 + len(data)}"
    return len(data)


def _sheet_months(wb):
    ws = wb.create_sheet("По месяцам")
    _title(ws, "По месяцам", "Выручка и количество оплат")
    _header(ws, 4, ["Месяц", "Оплат", "Выручка", "Средний чек"], [16, 12, 16, 16])
    rows = db.monthly_stats()
    data = [[r["month"], r["paid"], r["revenue"],
             round(r["revenue"] / r["paid"]) if r["paid"] else 0] for r in rows]
    end = _rows(ws, 5, data, {2: INT, 3: MONEY, 4: MONEY})
    ws.freeze_panes = "A5"
    if len(data) > 1:
        chart = BarChart()
        chart.title = "Выручка по месяцам"
        chart.height, chart.width = 8, 16
        chart.add_data(Reference(ws, min_col=3, min_row=4, max_row=end - 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=end - 1))
        ws.add_chart(chart, "F4")
    return len(data)


def _sheet_daily(wb):
    ws = wb.create_sheet("По дням")
    _title(ws, "По дням", "Динамика воронки")
    cols = ["Дата", "Пришли", "Согласия", "Открыли курс", "Оплат", "Выручка", "Дожимов"]
    _header(ws, 4, cols, [14, 12, 12, 16, 12, 16, 12])

    rows = db.daily_stats()
    data = [[_dt(r["day"], shift=False), r["starts"], r["consents"], r["product_views"],
             r["paid"], r["revenue"], r["drips"]] for r in rows]
    end = _rows(ws, 5, data, {1: DATE, 2: INT, 3: INT, 4: INT, 5: INT, 6: MONEY, 7: INT})
    ws.freeze_panes = "A5"

    if len(data) > 1:
        chart = LineChart()
        chart.title = "Динамика по дням"
        chart.height, chart.width = 9, 24
        chart.add_data(Reference(ws, min_col=2, max_col=7, min_row=4, max_row=end - 1),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=end - 1))
        ws.add_chart(chart, "I4")
    return len(data)


def _sheet_drips(wb):
    ws = wb.create_sheet("Дожимы")
    _title(ws, "Дожимы", "Кому и что ещё предстоит отправить")
    _header(ws, 4, ["ID", "Username", "Курс", "Отправлено шагов", "Следующий",
                    "Состояние"], [14, 18, 18, 18, 18, 26])

    rows = []
    for u in db.all_users():
        for d in db.user_drips(u["user_id"]):
            if d["muted"]:
                state = "отказался от напоминаний"
            elif d["stopped"]:
                state = "пауза (перешёл на другой курс или купил)"
            else:
                state = "идёт"
            rows.append([u["user_id"], f"@{u['username']}" if u["username"] else "",
                         SEG_TITLES.get(d["product"], d["product"]), d["step"],
                         _dt(d["next_at"]), state])
    _rows(ws, 5, rows, {1: INT, 4: INT, 5: DATE})
    ws.freeze_panes = "A5"
    return len(rows)


def _sheet_events(wb):
    ws = wb.create_sheet("События")
    _title(ws, "События", "Журнал действий, последние 5000")
    _header(ws, 4, ["Когда", "ID", "Событие", "Детали"], [20, 14, 22, 30])
    data = [[_dt(e["ts"]), e["user_id"], e["event"], e["meta"] or ""]
            for e in db.all_events()]
    _rows(ws, 5, data, {1: DATE, 2: INT})
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:D{4 + len(data)}"
    return len(data)


def _sheet_dashboard(wb, s):
    ws = wb.create_sheet("Дашборд", 0)
    _title(ws, "Бот «Код уважения» / «Пересборка»",
           f"Отчёт собран {datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M')} UTC")

    all_time = db.period_report(None)
    month = db.period_report(30)
    today = db.period_report(1)
    avg = round(all_time["revenue"] / all_time["paid"]) if all_time["paid"] else 0
    tiles = [
        ("Выручка всего", all_time["revenue"]), ("За 30 дней", month["revenue"]),
        ("Сегодня", today["revenue"]), ("Средний чек", avg),
        ("Оплат всего", all_time["paid"]), ("Всего в базе", s["total"]),
        ("Купили", s["buyers"]), ("Заблокировали бота", s["blocked"]),
    ]
    money_tiles = {0, 1, 2, 3}
    row = 4
    for i, (name, value) in enumerate(tiles):
        col = 1 + (i % 4) * 2
        r = row + (i // 4) * 3
        c1, c2 = ws.cell(row=r, column=col, value=name), ws.cell(row=r + 1, column=col, value=value)
        c1.font = Font(name=FONT, size=9, color="595959")
        c2.font = Font(name=FONT, bold=True, size=18, color=CLR_DARK)
        c2.number_format = MONEY if i in money_tiles else INT
        for cc in (c1, c2):
            cc.fill = PatternFill("solid", start_color=CLR_LIGHT)
            cc.alignment = Alignment(horizontal="center")
            ws.merge_cells(start_row=cc.row, start_column=col, end_row=cc.row, end_column=col + 1)
        ws.column_dimensions[get_column_letter(col)].width = 18
        ws.column_dimensions[get_column_letter(col + 1)].width = 6

    ws["A11"] = "Направление"; ws["B11"] = "Человек"
    seg = [("Пересборка", s["seg_per"]), ("Код уважения", s["seg_kod"]),
           ("Не выбрал", max(0, s["total"] - s["seg_per"] - s["seg_kod"]))]
    _rows(ws, 12, [list(x) for x in seg], {2: INT})

    ws["D11"] = "Источник"; ws["E11"] = "Человек"
    src = [("Instagram", s["src_ig"]), ("YouTube", s["src_yt"]), ("Прямой заход", s["src_direct"])]
    _rows(ws, 12, [[None, None, a, b] for a, b in src], {5: INT})
    for cell in ("A11", "B11", "D11", "E11"):
        ws[cell].font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        ws[cell].fill = PatternFill("solid", start_color=CLR_DARK)

    pie = PieChart()
    pie.title = "Направления"
    pie.height, pie.width = 8, 11
    pie.add_data(Reference(ws, min_col=2, min_row=11, max_row=14), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=12, max_row=14))
    ws.add_chart(pie, "A17")

    pie2 = PieChart()
    pie2.title = "Источники"
    pie2.height, pie2.width = 8, 11
    pie2.add_data(Reference(ws, min_col=5, min_row=11, max_row=14), titles_from_data=True)
    pie2.set_categories(Reference(ws, min_col=4, min_row=12, max_row=14))
    ws.add_chart(pie2, "H17")


def build_report():
    wb = Workbook()
    wb.remove(wb.active)

    s = db.stats()
    _sheet_funnel(wb)
    _sheet_payments(wb)
    _sheet_months(wb)
    _sheet_users(wb)
    _sheet_daily(wb)
    _sheet_drips(wb)
    _sheet_events(wb)
    _sheet_dashboard(wb, s)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"otchet-bot-{datetime.now(timezone.utc).strftime('%Y-%m-%d')}.xlsx"
    return name, buf
